"""Build a Book1.xlsx-style file for 78 projects in bulk_v2_out_khanh_78,
pulling data from the KHÁNH tab of source xlsx. D/E/F (Keywords/Headings/Descriptions) left empty.
"""
import csv
import re
import openpyxl

SOURCE_XLSX = r'C:\Users\Admin\Downloads\New folder\_source_khanh_trung.xlsx'
SOURCE_SHEET = 'KHÁNH'
INDEX_CSV = r'C:\Users\Admin\Documents\vippro campping\bulk_v2_out_khanh_78\_index.csv'
OUT_XLSX = r'C:\Users\Admin\Downloads\New folder\Book1_khanh_78.xlsx'

HEADERS = [
    'project_name', 'link_web', 'Final URL', 'Keywords', 'Headings',
    'Descriptions', 'Sitelinks', 'Callout Assets',
    'Top 1 Tier1', 'Top 2 Tier1', 'Top 3 Tier1', 'Top 4 Tier1', 'Top 5 Tier1',
    'Top 6 Tier1', 'Top 7 Tier1', 'Top 8 Tier1', 'Top 9 Tier1', 'Top 10 Tier1',
    ' CPC', 'Budget', 'Bid Strategy', 'Start Date', 'Networks',
    'EU Political Ads', 'Path 1', 'Path 2',
    'Tổng Volume Brand (Global)', 'Sum Top 10 Tier 1',
]


def normalize(name):
    if name is None:
        return ''
    n = str(name).lower().strip()
    return re.sub(r'[^a-z0-9]', '', n)


def fmt_val(v):
    if v is None:
        return None
    if isinstance(v, float) and v == int(v):
        return v
    return v


def main():
    wb_src = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    ws_src = wb_src[SOURCE_SHEET]

    # Build lookup dict
    sheet_idx = {}
    for r in range(2, ws_src.max_row + 1):
        name = ws_src.cell(r, 2).value
        if not name:
            continue
        key = normalize(name)
        if not key:
            continue
        if key not in sheet_idx:
            row = [ws_src.cell(r, c).value for c in range(1, 33)]
            sheet_idx[key] = (str(name).strip(), row)

    # Load 78-project index
    with open(INDEX_CSV, encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        idx_rows = [(row['Batch'], row['Project Name']) for row in reader]

    # Build output
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sheet1'
    ws.append(HEADERS)

    matched = 0
    unmatched = []
    for batch, idx_name in idx_rows:
        key = normalize(idx_name)
        hit = sheet_idx.get(key)
        if not hit:
            # try contains
            for sk, val in sheet_idx.items():
                if sk and len(sk) >= 4 and (sk in key or key in sk):
                    hit = val
                    break
        if hit:
            matched += 1
            sheet_name, r = hit
            # Map columns (1-indexed in original, here 0-indexed list of 32)
            row = [None] * 28
            row[0] = sheet_name                          # A project_name
            row[1] = r[2]                                # B link_web
            row[2] = r[5]                                # C Final URL (aff_link_text)
            row[3] = None                                # D Keywords
            row[4] = None                                # E Headings
            row[5] = None                                # F Descriptions
            row[6] = None                                # G Sitelinks
            row[7] = None                                # H Callout Assets
            for i in range(10):                          # I-R Top 1-10 Tier1
                row[8 + i] = fmt_val(r[10 + i])
            # S CPC: prefer THẦU (col 30, idx 29), fallback Max_CPC (idx 27)
            cpc = r[29] if r[29] not in (None, '') else r[27]
            row[18] = cpc
            row[19] = 100                                # T Budget
            row[20] = 'Manual CPC'                       # U Bid Strategy
            row[21] = None                               # V Start Date
            row[22] = 'Google search; Search partners'   # W Networks
            row[23] = 'No'                               # X EU Political Ads
            row[24] = None                               # Y Path 1
            row[25] = None                               # Z Path 2
            row[26] = fmt_val(r[9])                      # AA Tổng Volume Brand (Global)
            row[27] = fmt_val(r[20])                     # AB Sum Top 10 Tier 1
            ws.append(row)
        else:
            unmatched.append((batch, idx_name))
            ws.append([idx_name] + [None] * 27)

    wb.save(OUT_XLSX)
    print(f'Built: {OUT_XLSX}')
    print(f'Matched: {matched}/{len(idx_rows)}')
    if unmatched:
        print(f'Unmatched ({len(unmatched)}):')
        for b, n in unmatched:
            print(f'  - {b}: {n}')


if __name__ == '__main__':
    main()
